from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


@dataclass
class Section:
    text: str
    page: int | None = None


def parse_file(path: str) -> tuple[list[Section], str]:
    """Returns (sections, kind) where kind is 'markdown' or 'text'."""
    ext = Path(path).suffix.lower()

    if ext in {".md", ".markdown"}:
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        return [Section(text)], "markdown"

    if ext == ".txt":
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        return [Section(text)], "text"

    if ext == ".pdf":
        reader = PdfReader(path)
        return [Section(p.extract_text() or "", i + 1) for i, p in enumerate(reader.pages)], "text"

    if ext == ".docx":
        doc = Document(path)
        return [Section("\n".join(p.text for p in doc.paragraphs))], "text"

    if ext == ".pptx":
        prs = Presentation(path)
        out = []
        for i, slide in enumerate(prs.slides, 1):
            txt = "\n".join(
                s.text for s in slide.shapes if hasattr(s, "text") and s.text
            )
            out.append(Section(txt, i))
        return out, "text"

    raise ValueError(f"Unsupported file type: {ext}")
